#!/usr/bin/env python3
"""
comparar_documentos.py
Compara una factura PDF con un Excel de guías de envío y genera un reporte detallado.
"""

import argparse
import logging
import re
import subprocess
import sys
from pathlib import Path

# ── Auto-install dependencias ────────────────────────────────────────────────
def install(pkg):
    subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])

for _pkg in ["pymupdf", "pandas", "openpyxl", "rapidfuzz"]:
    try:
        __import__(_pkg if _pkg != "pymupdf" else "fitz")
    except ImportError:
        print(f"Instalando {_pkg}...")
        install(_pkg)

import fitz  # pymupdf — extractor PDF principal
_HAS_PDFPLUMBER = False  # deshabilitado: conflicto de cryptography en este entorno
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz

# ── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler("comparacion.log", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# ── Patrones comunes ─────────────────────────────────────────────────────────
GUIA_PATTERNS = [
    r'\b(\d{10,20})\b',
    r'(?:guia|guía|guide|tracking|#|no\.?)\s*:?\s*([A-Z0-9\-]{6,25})',
    r'\b([A-Z]{2,4}\d{8,15})\b',
]
MONTO_PATTERNS = [
    r'Q\.?\s*([\d,]+\.?\d*)',
    r'\$\s*([\d,]+\.?\d*)',
    r'([\d,]+\.\d{2})',
]
ESTADO_MAP = {
    'entregad': 'ENTREGADO',
    'deliver': 'ENTREGADO',
    'devuelt': 'DEVUELTO',
    'return': 'DEVUELTO',
    'transit': 'EN TRANSITO',
    'en ruta': 'EN TRANSITO',
    'pendient': 'PENDIENTE',
    'pending': 'PENDIENTE',
    'cancelad': 'CANCELADO',
    'cancel': 'CANCELADO',
    'retenid': 'RETENIDO',
    'held': 'RETENIDO',
}

# ── Helpers ──────────────────────────────────────────────────────────────────
def normalize_text(val):
    if val is None:
        return ""
    return str(val).strip().upper()


def normalize_amount(val):
    if val is None or val == "":
        return 0.0
    s = str(val).replace(",", "").replace("Q", "").replace("$", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def detect_estado(text):
    t = str(text).lower()
    for key, label in ESTADO_MAP.items():
        if key in t:
            return label
    return normalize_text(text) or "DESCONOCIDO"


def extract_guia(text):
    for pat in GUIA_PATTERNS:
        m = re.search(pat, str(text), re.IGNORECASE)
        if m:
            return m.group(1).strip().upper()
    return None


def extract_amount(text):
    for pat in MONTO_PATTERNS:
        m = re.search(pat, str(text).replace(" ", ""))
        if m:
            return normalize_amount(m.group(1))
    return None


def detect_currency(texts):
    combined = " ".join(str(t) for t in texts)
    if "Q" in combined:
        return "Q"
    if "$" in combined:
        return "$"
    return "Q"


# ── Extracción PDF ────────────────────────────────────────────────────────────
def extract_pdf_data(pdf_path: str) -> pd.DataFrame:
    log.info("Extrayendo datos del PDF: %s", pdf_path)
    records = []

    # --- Intentar con pdfplumber si está disponible ---
    try:
        if not _HAS_PDFPLUMBER:
            raise ImportError("pdfplumber no disponible")
        with pdfplumber.open(pdf_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                tables = page.extract_tables()
                for tbl in tables:
                    if not tbl:
                        continue
                    # Detectar fila de encabezado
                    header_row = 0
                    headers = [normalize_text(c) for c in tbl[0]]
                    log.info("Tabla pág %d — encabezados detectados: %s", page_num, headers)

                    col_guia = _find_col(headers, ["GUIA", "GUÍA", "TRACKING", "NO", "NUMERO", "N°", "GUIDE", "AWB", "CODIGO"])
                    col_estado = _find_col(headers, ["ESTADO", "STATUS", "ESTATUS", "CONDICION"])
                    col_costo = _find_col(headers, ["COSTO", "VALOR", "MONTO", "PRECIO", "IMPORTE", "TOTAL", "CARGO", "AMOUNT"])

                    for row in tbl[1:]:
                        if not row or all(c is None or str(c).strip() == "" for c in row):
                            continue
                        guia = _get_cell(row, col_guia)
                        estado = _get_cell(row, col_estado)
                        costo = _get_cell(row, col_costo)

                        # Fallback: buscar guía en toda la fila
                        if not guia:
                            for cell in row:
                                g = extract_guia(str(cell))
                                if g:
                                    guia = g
                                    break
                        if not costo:
                            for cell in row:
                                a = extract_amount(str(cell))
                                if a is not None:
                                    costo = a
                                    break

                        if guia:
                            records.append({
                                "guia": normalize_text(guia),
                                "estado_pdf": detect_estado(estado) if estado else "DESCONOCIDO",
                                "costo_pdf": normalize_amount(costo),
                            })

                # Si no hay tablas, extraer texto línea a línea
                if not tables:
                    text = page.extract_text() or ""
                    _parse_text_lines(text, records, "pdf")

    except Exception as e:
        log.warning("pdfplumber falló (%s), usando pymupdf como fallback", e)
        records = []
        doc = fitz.open(pdf_path)
        for page in doc:
            text = page.get_text()
            _parse_text_lines(text, records, "pdf")
        doc.close()

    if not records:
        log.warning("No se encontraron registros en el PDF con tablas; intentando extracción de texto plano.")
        records = []
        doc = fitz.open(pdf_path)
        for page in doc:
            text = page.get_text()
            _parse_text_lines(text, records, "pdf")
        doc.close()

    df = pd.DataFrame(records).drop_duplicates(subset=["guia"])
    log.info("PDF: %d registros extraídos", len(df))
    return df


def _find_col(headers, candidates):
    for c in candidates:
        for i, h in enumerate(headers):
            if c in h:
                return i
    return None


def _get_cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _parse_text_lines(text, records, source):
    lines = text.split("\n")
    for line in lines:
        guia = extract_guia(line)
        if not guia:
            continue
        estado_raw = ""
        for key in ESTADO_MAP:
            if key in line.lower():
                estado_raw = key
                break
        monto = extract_amount(line)
        records.append({
            "guia": normalize_text(guia),
            "estado_pdf": detect_estado(estado_raw) if estado_raw else "DESCONOCIDO",
            "costo_pdf": monto if monto is not None else 0.0,
        })


# ── Lectura Excel ─────────────────────────────────────────────────────────────
def read_excel_data(excel_path: str) -> pd.DataFrame:
    log.info("Leyendo Excel: %s", excel_path)
    xl = pd.ExcelFile(excel_path)
    sheet = xl.sheet_names[0]
    log.info("Hoja activa: %s", sheet)

    # Detectar fila de encabezado (buscar la primera fila con texto relevante)
    raw = pd.read_excel(excel_path, sheet_name=sheet, header=None, dtype=str)
    header_row = 0
    for i, row in raw.iterrows():
        vals = [normalize_text(v) for v in row if pd.notna(v)]
        if any(kw in v for v in vals for kw in ["GUIA", "TRACKING", "NO", "NUMERO", "ESTADO", "COSTO", "VALOR"]):
            header_row = i
            break

    df = pd.read_excel(excel_path, sheet_name=sheet, header=header_row, dtype=str)
    df.columns = [normalize_text(c) for c in df.columns]
    log.info("Columnas Excel: %s", list(df.columns))

    col_guia = _find_col_df(df.columns, ["GUIA", "GUÍA", "TRACKING", "NO", "NUMERO", "N°", "GUIDE", "AWB", "CODIGO"])
    col_estado = _find_col_df(df.columns, ["ESTADO", "STATUS", "ESTATUS", "CONDICION"])
    col_costo = _find_col_df(df.columns, ["COSTO", "VALOR", "MONTO", "PRECIO", "IMPORTE", "TOTAL", "CARGO", "AMOUNT"])

    log.info("Columna guía: %s | estado: %s | costo: %s", col_guia, col_estado, col_costo)

    records = []
    for _, row in df.iterrows():
        guia = normalize_text(row.get(col_guia, "")) if col_guia else None
        if not guia:
            # Buscar en toda la fila
            for val in row.values:
                g = extract_guia(str(val))
                if g:
                    guia = g
                    break
        if not guia or guia in ("NAN", "", "NONE"):
            continue

        estado = row.get(col_estado, "") if col_estado else ""
        costo_raw = row.get(col_costo, "") if col_costo else ""
        if not costo_raw or str(costo_raw).upper() in ("NAN", "NONE", ""):
            for val in row.values:
                a = extract_amount(str(val))
                if a is not None and a > 0:
                    costo_raw = a
                    break

        records.append({
            "guia": normalize_text(guia),
            "estado_excel": detect_estado(estado),
            "costo_excel": normalize_amount(costo_raw),
        })

    result = pd.DataFrame(records).drop_duplicates(subset=["guia"])
    log.info("Excel: %d registros leídos", len(result))
    return result


def _find_col_df(columns, candidates):
    for c in candidates:
        for col in columns:
            if c in col:
                return col
    return None


# ── Matching ──────────────────────────────────────────────────────────────────
def match_datasets(df_pdf: pd.DataFrame, df_excel: pd.DataFrame):
    pdf_set = set(df_pdf["guia"])
    excel_set = set(df_excel["guia"])

    both = pdf_set & excel_set
    only_excel = excel_set - pdf_set
    only_pdf = pdf_set - excel_set

    # Fuzzy match para guías no encontradas exactamente
    unmatched_pdf = only_pdf.copy()
    unmatched_excel = only_excel.copy()
    fuzzy_pairs = []
    for g_pdf in list(unmatched_pdf):
        for g_excel in list(unmatched_excel):
            score = fuzz.ratio(g_pdf, g_excel)
            if score >= 85:
                fuzzy_pairs.append((g_pdf, g_excel, score))
                log.info("Fuzzy match: %s <-> %s (score=%d)", g_pdf, g_excel, score)
                both.add(g_pdf)
                unmatched_pdf.discard(g_pdf)
                unmatched_excel.discard(g_excel)
                # Unificar guía en excel df
                df_excel.loc[df_excel["guia"] == g_excel, "guia"] = g_pdf
                break

    both_guias = list(both)
    only_excel_guias = list(unmatched_excel)
    only_pdf_guias = list(unmatched_pdf)

    merged = df_pdf[df_pdf["guia"].isin(both_guias)].merge(
        df_excel[df_excel["guia"].isin(both_guias)],
        on="guia", how="inner"
    )

    df_only_excel = df_excel[df_excel["guia"].isin(only_excel_guias)].copy()
    df_only_pdf = df_pdf[df_pdf["guia"].isin(only_pdf_guias)].copy()

    return merged, df_only_excel, df_only_pdf


# ── Estilos openpyxl ──────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9D9D9")
TOTAL_FONT = Font(bold=True)
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CURRENCY_FORMAT = '"Q"#,##0.00'


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_total_row(ws, row_idx):
    for cell in ws[row_idx]:
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = THIN_BORDER


def format_currency_cols(ws, col_indices, start_row=2):
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row):
        for i in col_indices:
            if i <= len(row):
                row[i - 1].number_format = CURRENCY_FORMAT


def autofit(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def apply_borders(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER


# ── Escritura de hojas ────────────────────────────────────────────────────────
def write_entregadas(wb, merged: pd.DataFrame, currency):
    ws = wb.create_sheet("Entregadas")
    entregadas = merged[
        (merged["estado_pdf"] == "ENTREGADO") & (merged["estado_excel"] == "ENTREGADO")
    ].copy()
    entregadas["diferencia"] = entregadas["costo_pdf"] - entregadas["costo_excel"]

    headers = ["N° Guía", "Estado", "Costo PDF", "Costo Excel", "Diferencia"]
    ws.append(headers)
    style_header(ws)

    for _, r in entregadas.iterrows():
        ws.append([
            r["guia"],
            r["estado_pdf"],
            r["costo_pdf"],
            r["costo_excel"],
            r["diferencia"],
        ])

    # Colorear diferencia
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=5, max_col=5):
        for cell in row:
            if cell.value == 0:
                cell.fill = GREEN_FILL
            elif cell.value is not None:
                cell.fill = RED_FILL

    # Totales
    total_row = [
        "TOTAL", "",
        entregadas["costo_pdf"].sum(),
        entregadas["costo_excel"].sum(),
        entregadas["diferencia"].sum(),
    ]
    ws.append(total_row)
    style_total_row(ws, ws.max_row)

    format_currency_cols(ws, [3, 4, 5])
    apply_borders(ws)
    autofit(ws)
    log.info("Hoja 'Entregadas': %d registros", len(entregadas))
    return len(entregadas), entregadas["costo_pdf"].sum(), entregadas["costo_excel"].sum()


def write_estado_diferente(wb, merged: pd.DataFrame, currency):
    ws = wb.create_sheet("Estado Diferente")
    diff = merged[
        ~((merged["estado_pdf"] == "ENTREGADO") & (merged["estado_excel"] == "ENTREGADO"))
    ].copy()

    headers = ["N° Guía", "Estado PDF", "Estado Excel", "Costo PDF", "Costo Excel"]
    ws.append(headers)
    style_header(ws)

    for _, r in diff.iterrows():
        ws.append([
            r["guia"],
            r["estado_pdf"],
            r["estado_excel"],
            r["costo_pdf"],
            r["costo_excel"],
        ])

    total_row = [
        "TOTAL", "", "",
        diff["costo_pdf"].sum(),
        diff["costo_excel"].sum(),
    ]
    ws.append(total_row)
    style_total_row(ws, ws.max_row)

    format_currency_cols(ws, [4, 5])
    apply_borders(ws)
    autofit(ws)
    log.info("Hoja 'Estado Diferente': %d registros", len(diff))
    return len(diff), diff["costo_pdf"].sum(), diff["costo_excel"].sum()


def write_solo_excel(wb, df: pd.DataFrame, currency):
    ws = wb.create_sheet("Solo en Excel")
    headers = ["N° Guía", "Estado", "Costo Excel"]
    ws.append(headers)
    style_header(ws)

    for _, r in df.iterrows():
        ws.append([r["guia"], r["estado_excel"], r["costo_excel"]])

    total_row = ["TOTAL", "", df["costo_excel"].sum()]
    ws.append(total_row)
    style_total_row(ws, ws.max_row)

    format_currency_cols(ws, [3])
    apply_borders(ws)
    autofit(ws)
    log.info("Hoja 'Solo en Excel': %d registros", len(df))
    return len(df), df["costo_excel"].sum()


def write_solo_pdf(wb, df: pd.DataFrame, currency):
    ws = wb.create_sheet("Solo en PDF")
    headers = ["N° Guía", "Estado", "Costo PDF"]
    ws.append(headers)
    style_header(ws)

    for _, r in df.iterrows():
        ws.append([r["guia"], r["estado_pdf"], r["costo_pdf"]])

    total_row = ["TOTAL", "", df["costo_pdf"].sum()]
    ws.append(total_row)
    style_total_row(ws, ws.max_row)

    format_currency_cols(ws, [3])
    apply_borders(ws)
    autofit(ws)
    log.info("Hoja 'Solo en PDF': %d registros", len(df))
    return len(df), df["costo_pdf"].sum()


def write_resumen(wb, stats: dict):
    ws = wb.create_sheet("Resumen General")
    headers = ["Categoría", "Cantidad", "Total Costo PDF", "Total Costo Excel"]
    ws.append(headers)
    style_header(ws)

    rows = [
        ("Entregadas",       stats["ent_cnt"],  stats["ent_pdf"],  stats["ent_excel"]),
        ("Estado Diferente", stats["dif_cnt"],  stats["dif_pdf"],  stats["dif_excel"]),
        ("Solo en Excel",    stats["xls_cnt"],  "-",               stats["xls_excel"]),
        ("Solo en PDF",      stats["pdf_cnt"],  stats["pdf_pdf"],  "-"),
    ]
    for r in rows:
        ws.append(list(r))

    # TOTAL GENERAL
    total_pdf = (
        (stats["ent_pdf"] if isinstance(stats["ent_pdf"], (int, float)) else 0) +
        (stats["dif_pdf"] if isinstance(stats["dif_pdf"], (int, float)) else 0) +
        (stats["pdf_pdf"] if isinstance(stats["pdf_pdf"], (int, float)) else 0)
    )
    total_excel = (
        (stats["ent_excel"] if isinstance(stats["ent_excel"], (int, float)) else 0) +
        (stats["dif_excel"] if isinstance(stats["dif_excel"], (int, float)) else 0) +
        (stats["xls_excel"] if isinstance(stats["xls_excel"], (int, float)) else 0)
    )
    total_cnt = stats["ent_cnt"] + stats["dif_cnt"] + stats["xls_cnt"] + stats["pdf_cnt"]
    ws.append(["TOTAL GENERAL", total_cnt, total_pdf, total_excel])
    style_total_row(ws, ws.max_row)

    # Formato moneda solo en celdas numéricas
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in [row[2], row[3]]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = CURRENCY_FORMAT

    apply_borders(ws)
    autofit(ws)


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Comparar factura PDF con Excel de guías")
    parser.add_argument("--pdf", required=True, help="Ruta al archivo PDF")
    parser.add_argument("--excel", required=True, help="Ruta al archivo Excel")
    parser.add_argument("--output", default="resultado_comparacion.xlsx", help="Archivo de salida")
    args = parser.parse_args()

    # Validar archivos
    for path, label in [(args.pdf, "PDF"), (args.excel, "Excel")]:
        if not Path(path).exists():
            log.error("Archivo %s no encontrado: %s", label, path)
            sys.exit(1)

    log.info("=" * 60)
    log.info("INICIO DE COMPARACIÓN")
    log.info("PDF: %s", args.pdf)
    log.info("Excel: %s", args.excel)
    log.info("=" * 60)

    df_pdf = extract_pdf_data(args.pdf)
    df_excel = read_excel_data(args.excel)

    if df_pdf.empty:
        log.error("No se pudo extraer ningún registro del PDF. Revisa la estructura del archivo.")
        sys.exit(1)
    if df_excel.empty:
        log.error("No se pudo leer ningún registro del Excel.")
        sys.exit(1)

    # Detectar moneda
    all_texts = list(df_pdf["costo_pdf"].astype(str)) + list(df_excel["costo_excel"].astype(str))
    currency = detect_currency(all_texts)
    log.info("Moneda detectada: %s", currency)

    merged, df_only_excel, df_only_pdf = match_datasets(df_pdf, df_excel)

    # Crear workbook
    wb = Workbook()
    wb.remove(wb.active)  # quitar hoja por defecto

    ent_cnt, ent_pdf, ent_excel = write_entregadas(wb, merged, currency)
    dif_cnt, dif_pdf, dif_excel = write_estado_diferente(wb, merged, currency)
    xls_cnt, xls_excel = write_solo_excel(wb, df_only_excel, currency)
    pdf_cnt, pdf_pdf = write_solo_pdf(wb, df_only_pdf, currency)

    write_resumen(wb, {
        "ent_cnt": ent_cnt, "ent_pdf": ent_pdf, "ent_excel": ent_excel,
        "dif_cnt": dif_cnt, "dif_pdf": dif_pdf, "dif_excel": dif_excel,
        "xls_cnt": xls_cnt, "xls_excel": xls_excel,
        "pdf_cnt": pdf_cnt, "pdf_pdf": pdf_pdf,
    })

    wb.save(args.output)
    log.info("Archivo guardado: %s", args.output)

    # Resumen en consola
    print("\n" + "=" * 60)
    print("RESUMEN DE COMPARACIÓN")
    print("=" * 60)
    print(f"  Entregadas (en ambos):    {ent_cnt:>5}  PDF={ent_pdf:>12,.2f}  Excel={ent_excel:>12,.2f}")
    print(f"  Estado diferente:         {dif_cnt:>5}  PDF={dif_pdf:>12,.2f}  Excel={dif_excel:>12,.2f}")
    print(f"  Solo en Excel:            {xls_cnt:>5}  Excel={xls_excel:>12,.2f}")
    print(f"  Solo en PDF:              {pdf_cnt:>5}  PDF={pdf_pdf:>12,.2f}")
    print(f"\n  Resultado guardado en: {args.output}")
    print("=" * 60)
    log.info("COMPARACIÓN FINALIZADA")


if __name__ == "__main__":
    main()
